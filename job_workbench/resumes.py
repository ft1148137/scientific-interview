"""Resume imports retain originals; exports cite immutable, confirmed experience versions."""
from copy import deepcopy
from difflib import unified_diff
from html import escape
from pathlib import Path
import re
import shutil
from uuid import uuid4
from .contracts import require, identifier
from .debriefs import validate_debrief_refs
from .resources import FONT_FILE
from .store import current, digest, now, add_event, WorkbenchError, atomic_json, data_path


def import_resume(store, path):
    path = Path(path)
    suffix = path.suffix.lower()
    require(suffix in {".pdf", ".docx", ".txt", ".md"}, "支持 PDF、DOCX、TXT、Markdown")
    raw = path.read_bytes()
    require(len(raw) <= 25 * 1024 * 1024, "简历超过 25 MiB")
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise WorkbenchError('请安装简历依赖: pip install -e ".[resume]"') from exc
        reader = PdfReader(path)
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise WorkbenchError('请安装简历依赖: pip install -e ".[resume]"') from exc
        doc = Document(path)
        text = "\n".join([p.text for p in doc.paragraphs] + [" | ".join(c.text for c in r.cells) for t in doc.tables for r in t.rows])
    else:
        text = raw.decode("utf-8-sig")
    original = store.blob(raw, suffix)
    imported = {"id": digest(raw.hex())[:20], "original": original, "text": text, "at": now(), "status": "needs_review" if text.strip() else "needs_transcription", "note": "这是导入原件 ID，用 resume-draft --source-import 引用；--parent 只接受已保存的 resume-… 版本。先由 AI 从原文整理 profiles/personal 和 experiences，用户核对后再确认，不要求用户重复提供已有事实。"}
    store.update(lambda state: state["imports"].setdefault(imported["id"], imported))
    return imported


def draft_resume(state, target, experience_ids=None, parent_id=None, source_import_id=None, rebuild=False):
    require(parent_id is None or parent_id in state["resumes"], "--parent 需要已保存的 resume-… 版本；导入原件请改用 --source-import <导入ID>。先 show resumes 查看可用版本。")
    if parent_id and source_import_id is None:
        source_import_id = state["resumes"][parent_id]["draft"].get("source_import_id")
    require(source_import_id is None or source_import_id in state["imports"], "导入原件不存在，请先 resume-import 或 show imports")
    if parent_id and not rebuild:
        require(experience_ids is None, "续改父稿保留原经历选择；如需从经历库重建，请显式加 --rebuild 并指定 --experience。也可直接编辑续改草稿的条目取舍。")
        draft = deepcopy(state["resumes"][parent_id]["draft"])
        draft.update(target=target, parent_id=parent_id, source_import_id=source_import_id)
        return draft
    profile = current(state["profiles"], "personal")
    require(profile and profile["confirmed"], "先从用户提供的资料整理基础信息，再用 library-put profiles personal <资料.json> --confirmed 保存；已有记录需加当前 --expected-version。")
    ids = experience_ids if experience_ids is not None else list(state["experiences"])
    entries = []
    for key in ids:
        entry = current(state["experiences"], key)
        require(entry and entry["confirmed"], f"经历 {key} 尚未确认")
        data = entry["data"]
        if data.get("deleted"):
            continue
        bullets = [data[field] for field in ("responsibility", "actions", "result") if data.get(field)]
        if bullets:
            item = {"experience_id": key, "experience_version": entry["version"], "title": data["title"], "organization": data.get("organization", ""), "dates": data.get("dates", ""), "bullets": bullets}
            if data.get("facts"):
                item["fact_context"] = deepcopy(data["facts"])
            entries.append(item)
    require(entries, "至少需要一条已确认的职责、行动或成果")
    return {"target": target, "profile": deepcopy(profile["data"]), "profile_version": profile["version"], "entries": entries, "parent_id": parent_id, "source_import_id": source_import_id, "change_notes": ["从已确认的经历生成，可按目标岗位调整取舍与表达。"]}


def source_updates(state, draft):
    refs = [("profiles", "personal", draft["profile_version"])] + [("experiences", e["experience_id"], e["experience_version"]) for e in draft["entries"]]
    refs += [("debriefs", r["id"], r["version"]) for r in draft.get("debrief_refs", [])]
    updates = []
    for collection, key, version in refs:
        latest = current(state.get(collection, {}), key)
        if latest and latest["version"] != version:
            updates.append({"collection": collection, "id": key, "used_version": version, "latest_version": latest["version"], "confirmed": latest["confirmed"], "deleted": bool(latest["data"].get("deleted"))})
    return updates


def save_resume(state, draft, confirmed):
    require(confirmed, "简历改写需用户确认事实，使用 --confirmed 保存最终版本")
    require(isinstance(draft, dict) and draft.get("target") and draft.get("entries"), "简历需要目标方向和经历条目")
    profile_versions = state["profiles"].get("personal", [])
    pv = draft.get("profile_version")
    require(type(pv) is int and 1 <= pv <= len(profile_versions), "基础资料版本不存在")
    require(profile_versions[pv-1]["confirmed"] and draft.get("profile") == profile_versions[pv-1]["data"], "基础资料需与确认版本一致，先更新资料再生成简历")
    for entry in draft["entries"]:
        history = state["experiences"].get(entry.get("experience_id"), [])
        version = entry.get("experience_version")
        require(type(version) is int and 1 <= version <= len(history) and history[version - 1]["confirmed"], "简历引用了不存在或未确认的经历")
        require(isinstance(entry.get("bullets"), list) and entry["bullets"] and all(isinstance(b, str) and b.strip() for b in entry["bullets"]), "经历条目不能为空")
        # Wording may change after confirmation; factual headers must stay attached to their source.
        source = history[version-1]["data"]
        for field in ("title", "organization", "dates"):
            require(entry.get(field, "") == source.get(field, ""), f"{field} 与经历版本不一致，请先修改经历库")
        require(entry.get("fact_context", {}) == source.get("facts", {}), "事实限定与经历版本不一致；不能在简历中删除或更改阶段、指标状态，请先核对经历库。")
    parent = draft.get("parent_id")
    require(parent is None or parent in state["resumes"], "父版本不存在；parent_id 只接受已保存简历，导入原件应放入 source_import_id")
    require(draft.get("source_import_id") is None or draft["source_import_id"] in state["imports"], "引用的导入原件不存在")
    validate_debrief_refs(state, draft)
    resume_id = "resume-" + digest(draft)[:20]
    if resume_id in state["resumes"]:
        return state["resumes"][resume_id]
    record = {"id": resume_id, "at": now(), "confirmed": True, "draft": deepcopy(draft)}
    state["resumes"][resume_id] = record
    add_event(state, "resume_created", resume_id=resume_id, parent_id=parent)
    return record


def resume_diff(state, record):
    """Compare displayed facts and wording, independently of user-written notes."""
    parent = state["resumes"].get(record["draft"].get("parent_id"))
    if not parent:
        return ""
    def lines(draft):
        result = ["目标：" + draft["target"]]
        labels = {"name": "姓名", "email": "邮箱", "phone": "电话", "education": "教育", "skills": "技能", "summary": "简介", "location": "所在地"}
        result += [f"{labels.get(key, key)}：{value}" for key, value in sorted(draft["profile"].items())]
        for entry in draft["entries"]:
            result += ["经历：" + " · ".join(entry.get(k, "") for k in ("title", "organization", "dates"))]
            result += ["- " + bullet for bullet in entry["bullets"]]
            result += ["范围：" + line for line in fact_lines(entry.get("fact_context", {}))]
        return result
    return "\n".join(unified_diff(lines(parent["draft"]), lines(record["draft"]), fromfile="上一版本", tofile="当前版本", lineterm="")) or "内容未变化，仅版本说明等元数据有调整。"


def fact_lines(facts):
    lines = [f"{label}：{facts[key]}" for key, label in (("stage", "项目阶段"), ("responsibility_boundary", "职责范围")) if facts.get(key)]
    labels = {"target": "目标值", "forecast": "预测值", "actual": "已实现", "settled": "已结算", "unknown": "状态待核验"}
    for metric in facts.get("metrics", []):
        if metric.get("disclosure") != "public":
            continue
        details = [labels[metric["status"]], metric.get("value"), metric.get("unit"), metric.get("period"), metric.get("definition"), metric.get("owner_scope")]
        lines.append(metric["name"] + "：" + " · ".join(x for x in details if x))
    return lines


def credential_lines(profile):
    kinds = {"training": "培训", "certificate": "证书", "authorization": "授权"}
    statuses = {"held": "已持有", "expired": "已过期", "pending": "待取得", "unknown": "状态待核验"}
    return [" · ".join(x for x in (kinds[c["kind"]] + "：" + c["name"], statuses[c["status"]], c.get("issuer"), c.get("date"), "有效期至 " + c["expires_on"] if c.get("expires_on") else "", "注册状态：" + c["registration"] if c.get("registration") else "") if x)
            for c in profile.get("credentials", []) if c.get("include_in_resume")]


def embed_resume_font(doc):
    """Keep editable Chinese text readable without installing a system font."""
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import qn
    from docx.opc.part import Part
    from docx.opc.packuri import PackURI
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.oxml import serialize_part_xml
    font_key = uuid4()
    key = font_key.bytes[::-1]
    raw = bytearray(FONT_FILE.read_bytes())
    # OOXML font embedding obfuscates the first 32 bytes with the reversed GUID.
    for i in range(32):
        raw[i] ^= key[i % 16]
    table = doc.part.part_related_by(RT.FONT_TABLE)
    root = parse_xml(table.blob)
    font = OxmlElement("w:font")
    font.set(qn("w:name"), "Noto Sans SC")
    part = Part(PackURI("/word/fonts/NotoSansSC-Regular.odttf"), "application/vnd.openxmlformats-officedocument.obfuscatedFont", bytes(raw), doc.part.package)
    embedded = OxmlElement("w:embedRegular")
    embedded.set(qn("r:id"), table.relate_to(part, RT.FONT))
    embedded.set(qn("w:fontKey"), "{" + str(font_key).upper() + "}")
    font.append(embedded)
    root.append(font)
    table._blob = serialize_part_xml(root)
    doc.settings.element.append(OxmlElement("w:embedTrueTypeFonts"))
    doc.settings.element.append(OxmlElement("w:embedSystemFonts"))


def export_resume(store, resume_id, output=None):
    state = store.read()
    require(resume_id in state["resumes"], "简历版本不存在")
    record = state["resumes"][resume_id]
    draft = record["draft"]
    target = data_path(output if output else store.root / "exports" / resume_id)
    canonical = data_path(store.root / "exports" / resume_id)
    files = ["resume.docx", "resume.pdf", "resume.md", "changes.md", "resume.json"]
    for directory in (target, canonical):
        for filename in files:
            data_path(directory / filename)
    target.mkdir(parents=True, exist_ok=True)
    profile = draft["profile"]
    lines = ["# " + profile.get("name", ""), draft["target"], " · ".join(str(profile[k]) for k in ("email", "phone", "location") if profile.get(k))]
    for field in ("summary", "education", "skills"):
        if profile.get(field):
            lines += ["", str(profile[field])]
    lines += credential_lines(profile)
    for entry in draft["entries"]:
        lines += ["", "## " + " · ".join(x for x in (entry["title"], entry.get("organization"), entry.get("dates")) if x)]
        lines += ["- " + b for b in entry["bullets"]]
        lines += fact_lines(entry.get("fact_context", {}))
    text = "\n".join(lines) + "\n"
    (target / "resume.md").write_text(text, encoding="utf-8")
    atomic_json(target / "resume.json", record)
    changes = "# 修改说明\n\n" + "\n".join(draft.get("change_notes", []))
    if draft.get("parent_id"):
        changes += "\n\n## 系统比较（与父版本）\n\n```diff\n" + resume_diff(state, record) + "\n```"
    (target / "changes.md").write_text(changes + "\n", encoding="utf-8")
    try:
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        from docx.oxml.ns import qn
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    except ImportError as exc:
        raise WorkbenchError('Markdown 已保存；Word/PDF 需 pip install -e ".[resume]"') from exc
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.65)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Noto Sans SC"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Noto Sans SC")
    for style in doc.styles:
        for border in list(style._element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
        if style.type == 1:
            style.font.color.rgb = RGBColor(0, 0, 0)
        for fonts in style._element.iter(qn("w:rFonts")):
            for attr in list(fonts.attrib):
                if "theme" in attr.lower():
                    del fonts.attrib[attr]
            fonts.set(qn("w:eastAsia"), "Noto Sans SC")
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Heading 2"].font.size = Pt(12)
    doc.styles["Normal"].font.size = Pt(10)
    pdfmetrics.registerFont(TTFont("ResumeCN", str(FONT_FILE)))
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("CN", fontName="ResumeCN", fontSize=10, leading=16, spaceAfter=5, wordWrap="CJK"))
    styles.add(ParagraphStyle("CNTitle", parent=styles["CN"], fontSize=22, leading=28, spaceAfter=12))
    styles.add(ParagraphStyle("CNHeading", parent=styles["CN"], fontSize=12, leading=18, spaceBefore=9, spaceAfter=5))
    story = []
    for line in lines:
        if not line:
            continue
        if line.startswith("# "):
            content, style, ps = line[2:], "Title", "CNTitle"
        elif line.startswith("## "):
            content, style, ps = line[3:], "Heading 2", "CNHeading"
        else:
            content, style, ps = line, "Normal", "CN"
        paragraph = doc.add_paragraph(content, style)
        for run in paragraph.runs:
            fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
            for script in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn("w:" + script), "Noto Sans SC")
        story.append(Paragraph(escape(content).replace("\n", "<br/>"), styles[ps]))
    embed_resume_font(doc)
    doc.save(target / "resume.docx")
    SimpleDocTemplate(str(target / "resume.pdf"), topMargin=42, bottomMargin=42, leftMargin=45, rightMargin=45).build(story)
    if target.resolve() != canonical.resolve():
        canonical.mkdir(parents=True, exist_ok=True)
        for filename in files:
            shutil.copy2(target / filename, canonical / filename)
    return {"id": resume_id, "directory": str(target.resolve()), "files": files,
            "paths": {name: str((target / name).resolve()) for name in files},
            "workbench_directory": str(canonical), "next_action": "打开 PDF/Word 核对全部页面；DOCX 已嵌入中文字体（约 6 MB），仍需在实际使用的办公软件核验。运行 render 后，工作台可下载此版本。"}
